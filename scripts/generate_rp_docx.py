import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Set subtle table borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def build_research_proposal():
    doc = docx.Document()

    # Page Setup: Standard A4, 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    # Styles definition
    COLOR_PRIMARY = RGBColor(26, 54, 93)      # Deep Navy #1A365D
    COLOR_SECONDARY = RGBColor(43, 108, 176)  # Slate Blue #2B6CB0
    COLOR_TEXT = RGBColor(45, 55, 72)         # Dark Gray #2D3748

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # 1. Header / Meta Box
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("RESEARCH PROPOSAL")
    t_run.font.name = 'Calibri'
    t_run.font.size = Pt(13)
    t_run.font.bold = True
    t_run.font.color.rgb = COLOR_SECONDARY

    main_title = doc.add_paragraph()
    main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    main_title.paragraph_format.space_after = Pt(12)
    mt_run = main_title.add_run("Decentralized Multi-Agent Coordination and Closed-Loop Evolution for Predictive Maintenance in Critical Infrastructure")
    mt_run.font.name = 'Georgia'
    mt_run.font.size = Pt(17)
    mt_run.font.bold = True
    mt_run.font.color.rgb = COLOR_PRIMARY

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Target Degree:", "Ph.D. / M.Sc. by Research in Computer Science"),
        ("Target Supervisor:", "Professor Gregory O'Hare (Head of School, Professor of Artificial Intelligence)"),
        ("Target Institution:", "School of Computer Science and Statistics, Trinity College Dublin (TCD)"),
        ("Applicant:", "Li Han (M.Sc. NTU GPA 4.0/5.0, B.Sc. First Class Honours QMUL)")
    ]

    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        row.cells[0].width = Inches(1.8)
        row.cells[1].width = Inches(4.5)

        pk = row.cells[0].paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.font.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = COLOR_PRIMARY

        pv = row.cells[1].paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = COLOR_TEXT

        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_background(row.cells[1], "F7FAFC")
        set_cell_margins(row.cells[0], top=60, bottom=60, left=100, right=100)
        set_cell_margins(row.cells[1], top=60, bottom=60, left=100, right=100)

    set_table_borders(meta_table, color="CBD5E0", sz="4")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def add_section_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title_text)
        run.font.name = 'Georgia'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY

        # Bottom divider border under heading
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="2B6CB0"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
        return p

    def add_sub_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    # 1. Abstract
    add_section_heading("1. Abstract")
    p_abs = doc.add_paragraph()
    p_abs.paragraph_format.space_after = Pt(6)
    p_abs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.add_run(
        "Modern critical infrastructure—spanning energy utility networks, smart grids, and hyperscale data center clusters—exhibits complex cyber-physical interdependencies where failures in physical environments (power, thermal dynamics, mechanical telemetry) rapidly propagate into digital workloads. Traditional operations remain severely bottlenecked by siloed diagnostic data, manual triage latency, and the vulnerability of naive Large Language Model (LLM) agents to context explosion and tool hallucination. This research proposal investigates an autonomous, decentralized Multi-Agent System (MAS) architecture engineered specifically for predictive maintenance, early fault warning, and trustworthy automated remediation in critical infrastructure. We propose an asynchronous Star-Topology Blackboard state bus with dynamic context pruning to constrain inter-agent communication overhead, coupled with deterministic Abstract Syntax Tree (AST) guardrails and a chaos-injected Direct Preference Optimization (DPO) data flywheel. The overarching objective is to achieve sub-second cross-domain root-cause localization, eliminate tool hallucination, and enforce zero-trust safety constraints for automated incident remediation."
    )

    # 2. Motivation
    add_section_heading("2. Research Motivation & Background")
    p_mot1 = doc.add_paragraph()
    p_mot1.paragraph_format.space_after = Pt(6)
    p_mot1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_mot1.add_run(
        "In mission-critical infrastructure, operational failures incur massive financial loss, service downtime, and safety hazards. Addressing these challenges through autonomous artificial intelligence faces two primary technical barriers:"
    )

    p_b1 = doc.add_paragraph(style='List Bullet')
    p_b1.paragraph_format.space_after = Pt(3)
    r_b1_title = p_b1.add_run("The Cyber-Physical Telemetry Divide: ")
    r_b1_title.bold = True
    p_b1.add_run(
        "Physical telemetry (L1–L2: power distribution units, cooling systems, environmental sensors) and digital infrastructure (L3–L7: hypervisors, containerized microservices, distributed storage) operate in isolated silos. When localized physical anomalies occur, cross-domain root-cause correlation currently relies on slow manual inspection, easily missing the critical 10-minute remediation window."
    )

    p_b2 = doc.add_paragraph(style='List Bullet')
    p_b2.paragraph_format.space_after = Pt(6)
    r_b2_title = p_b2.add_run("Non-Deterministic Multi-Agent Failure Modes: ")
    r_b2_title.bold = True
    p_b2.add_run(
        "In agentic systems, mesh peer-to-peer communication results in quadratic O(N²) message inflation, rapid context dilution, and cascading hallucinations. Moreover, executing mutating commands (e.g., node eviction, circuit switching) on live infrastructure without deterministic verification poses existential operational risks."
    )

    p_mot2 = doc.add_paragraph()
    p_mot2.paragraph_format.space_after = Pt(6)
    p_mot2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_mot2.add_run(
        "Professor Gregory O'Hare's internationally recognized leadership in Distributed Artificial Intelligence (DAI), Multi-Agent Systems (MAS), and ubiquitous computing—especially his recently funded €865k research initiative with CKDelta on AI-driven predictive maintenance for utility networks—provides the ideal theoretical and applied environment to resolve these critical challenges."
    )

    # 3. Research Questions & Objectives
    add_section_heading("3. Research Questions & Core Objectives")

    rq_table = doc.add_table(rows=3, cols=2)
    rq_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rq_table.autofit = False

    rqs = [
        ("RQ1: Decentralized State Governance", "How can we architect an asynchronous multi-agent coordination protocol that minimizes message complexity and context degradation while correlating high-frequency physical telemetry with distributed workload traces?"),
        ("RQ2: Deterministic Safety & Guardrails", "How can we establish formal AST-level parameter validation and cryptographic approval gates within the Model Context Protocol (MCP) to mathematically eliminate tool hallucination and prevent unauthorized mutations on critical physical assets?"),
        ("RQ3: Closed-Loop Evolution via DPO Flywheel", "How can automated chaos simulation sandboxes and execution telemetry be harnessed to curate contrastive DPO datasets that continuously optimize compact domain models for zero-shot operational precision?")
    ]

    for idx, (title, desc) in enumerate(rqs):
        row = rq_table.rows[idx]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.1)

        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(title)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(desc)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_TEXT

        bg = "F7FAFC" if idx % 2 == 0 else "EDF2F7"
        set_cell_background(row.cells[0], bg)
        set_cell_background(row.cells[1], bg)
        set_cell_margins(row.cells[0], top=80, bottom=80, left=100, right=100)
        set_cell_margins(row.cells[1], top=80, bottom=80, left=100, right=100)

    set_table_borders(rq_table, color="CBD5E0", sz="4")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Methodology
    add_section_heading("4. Proposed Methodology & System Architecture")
    p_meth_intro = doc.add_paragraph()
    p_meth_intro.paragraph_format.space_after = Pt(6)
    p_meth_intro.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_meth_intro.add_run(
        "We structure the technical research program into three synergistic work packages spanning multi-agent topology, deterministic verification gates, and continuous preference optimization:"
    )

    add_sub_heading("Work Package 1: Star-Topology Blackboard Bus & Dynamic Context Pruning (Year 1)")
    p_wp1 = doc.add_paragraph()
    p_wp1.paragraph_format.space_after = Pt(4)
    p_wp1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_wp1.add_run(
        "To eliminate inter-agent chatter and token blowup, we construct an asynchronous Star-Topology Blackboard. Specialized domain workers (Physical Facilities Agent, Cloud Compute Agent, Security Isolation Agent) interact strictly via structured evidence items (EvidenceItem: source_agent, target_entity, time_to_failure, confidence_score, distilled_evidence). A dynamic Context Pruning Engine analyzes downstream task dependencies to deliver strictly filtered context slices, achieving a >60% token reduction and eliminating hallucination propagation."
    )

    add_sub_heading("Work Package 2: Dual-Core AST Evaluation Engine & Zero-Trust Sandboxes (Year 2)")
    p_wp2 = doc.add_paragraph()
    p_wp2.paragraph_format.space_after = Pt(4)
    p_wp2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_wp2.add_run(
        "We implement formal Abstract Syntax Tree (AST) validation algorithms integrated into dynamic Model Context Protocol (MCP) gateways. Every tool call generated by an LLM agent undergoes static type checking, range validation, and schema bounds enforcement before execution. For mutating operations (e.g., node drain, circuit switching), the gate enforces an ephemeral cryptographic token requiring supervisor approval, completely preventing unauthorized destructive actions on live infrastructure."
    )

    add_sub_heading("Work Package 3: Chaos-Injected Evaluation & DPO Self-Evolution Flywheel (Year 3–4)")
    p_wp3 = doc.add_paragraph()
    p_wp3.paragraph_format.space_after = Pt(6)
    p_wp3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_wp3.add_run(
        "We construct an automated chaos sandbox that injects simulated network jitter, partial sensor drops, and malformed telemetry streams into agent evaluation suites. Trajectories and scores from a dual-core evaluator (Deterministic Engineering Metrics + LLM-as-a-Judge) are persisted into an embedded database, automatically extracting positive and negative preference pairs. These datasets drive continuous Direct Preference Optimization (DPO) fine-tuning of compact domain models (7B–14B parameters) via DeepSpeed and Ray, establishing a self-improving operational intelligence loop."
    )

    # 5. Work Plan Table
    add_section_heading("5. 4-Year PhD Work Plan & Milestones")

    plan_table = doc.add_table(rows=5, cols=3)
    plan_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    plan_table.autofit = False

    headers = ["Timeline", "Core Research Focus & Milestones", "Key Deliverables & Target Publications"]
    hdr_row = plan_table.rows[0]
    for i, h in enumerate(headers):
        hdr_row.cells[i].paragraphs[0].add_run(h).bold = True
        hdr_row.cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr_row.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_row.cells[i], "1A365D")
        set_cell_margins(hdr_row.cells[i], top=80, bottom=80, left=100, right=100)
    hdr_row.cells[0].width = Inches(1.1)
    hdr_row.cells[1].width = Inches(2.9)
    hdr_row.cells[2].width = Inches(2.3)

    plan_rows = [
        ("Year 1", "Theoretical modeling of Star Blackboard topology; Context pruning algorithms; Prototype integration with sensor telemetry streams.", "Literature review paper; AAMAS / IJCAI submission on Blackboard State Governance."),
        ("Year 2", "AST schema validation engine; Dynamic MCP tool routing; Cryptographic approval gate implementation and sandbox safety verification.", "Conference paper on Trustworthy Agent Execution; IEEE TKDE / ACM TAAS journal submission."),
        ("Year 3", "Chaos sandbox testbed construction; SQLite telemetry data flywheel; Distributed DPO fine-tuning using DeepSpeed and Ray on 7B/14B SLMs.", "Benchmark release: Critical Infrastructure Chaos Eval; Top AI Conference publication (NeurIPS/AAAI)."),
        ("Year 4", "System deployment across large-scale physical/digital testbeds; Comprehensive empirical evaluation; PhD thesis synthesis and defense.", "Doctoral Dissertation; SFI ADAPT / CKDelta Project Final Report; Journal publications.")
    ]

    for idx, (time_col, focus_col, deliv_col) in enumerate(plan_rows, start=1):
        row = plan_table.rows[idx]
        row.cells[0].width = Inches(1.1)
        row.cells[1].width = Inches(2.9)
        row.cells[2].width = Inches(2.3)

        row.cells[0].paragraphs[0].add_run(time_col).bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(9)
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = COLOR_PRIMARY

        row.cells[1].paragraphs[0].add_run(focus_col).font.size = Pt(9)
        row.cells[2].paragraphs[0].add_run(deliv_col).font.size = Pt(9)

        bg = "FFFFFF" if idx % 2 != 0 else "F7FAFC"
        set_cell_background(row.cells[0], bg)
        set_cell_background(row.cells[1], bg)
        set_cell_background(row.cells[2], bg)
        set_cell_margins(row.cells[0], top=60, bottom=60, left=80, right=80)
        set_cell_margins(row.cells[1], top=60, bottom=60, left=80, right=80)
        set_cell_margins(row.cells[2], top=60, bottom=60, left=80, right=80)

    set_table_borders(plan_table, color="CBD5E0", sz="4")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 6. Expected Contributions & Alignment
    add_section_heading("6. Alignment with Professor Gregory O'Hare's Laboratory")
    p_align = doc.add_paragraph()
    p_align.paragraph_format.space_after = Pt(6)
    p_align.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_align.add_run(
        "This proposal directly complements Professor O'Hare's €865,000 research collaboration with CKDelta focusing on AI-driven predictive maintenance and early warning systems for the utilities and critical infrastructure sectors. The applicant brings a distinguished academic foundation—an M.Sc. in Computer Control & Automation from Nanyang Technological University (GPA: 4.0/5.0) and a B.Sc. with First Class Honours in Electrical Engineering from Queen Mary University of London—coupled with hands-on production experience architecting enterprise multi-agent platforms and distributed LLM fine-tuning pipelines. This unique synthesis of control engineering, distributed systems, and modern agentic AI ensures immediate productivity, rigorous research execution, and high-impact contributions to the research group at Trinity College Dublin and the SFI ADAPT Centre."
    )

    # 7. Key References
    add_section_heading("7. Selected Key References")
    refs = [
        "O'Hare, G. M. P., & Jennings, N. R. (Eds.). Foundations of Distributed Artificial Intelligence. John Wiley & Sons.",
        "Wooldridge, M. (2009). An Introduction to MultiAgent Systems (2nd ed.). John Wiley & Sons.",
        "Rafailov, R., Sharma, A., Mitchell, E., Ermon, S., Manning, C. D., & Finn, C. (2023). Direct Preference Optimization: Your Language Model is Secretly a Reward Model. Advances in Neural Information Processing Systems (NeurIPS 2023).",
        "Schick, T., Dwivedi-Yu, J., Dessì, R., Raileanu, R., Lomeli, M., Zettlemoyer, L., Cancedda, N., & Scialom, T. (2023). Toolformer: Language Models Can Teach Themselves to Use Tools. Advances in Neural Information Processing Systems (NeurIPS 2023).",
        "SFI ADAPT / INSIGHT Centre for Data Analytics Technical Reports on Ubiquitous Computing and Critical Infrastructure Resilience (2023–2025)."
    ]
    for r in refs:
        p_r = doc.add_paragraph(style='List Bullet')
        p_r.paragraph_format.space_after = Pt(2)
        r_run = p_r.add_run(r)
        r_run.font.size = Pt(9)
        r_run.font.color.rgb = COLOR_TEXT

    # Output paths
    output_dir = "/Users/hanli/Downloads/Jinjing_Agent/docs"
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, "Research_Proposal_Li_Han_TCD.docx")
    doc.save(out_path)

    # Also save to root for easy access
    doc.save("/Users/hanli/Downloads/Jinjing_Agent/Research_Proposal_Li_Han_TCD.docx")
    print(f"Successfully generated Word RP at: {out_path}")

if __name__ == "__main__":
    build_research_proposal()
